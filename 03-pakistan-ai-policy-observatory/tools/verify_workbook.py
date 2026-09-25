#!/usr/bin/env python3
"""Independently verify the Pakistan AI Policy Monitor workbook."""

from __future__ import annotations

import argparse
import csv
import math
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook


__all__ = ["verify_workbook"]

ROOT = Path(__file__).resolve().parents[1]
INTERMEDIATE = ROOT / "intermediate"
OUTPUT_DIR = ROOT / "outputs" / "01a0653c-96d7-7c92-918b-3975d6716164"
WORKBOOK_PATH = OUTPUT_DIR / (
    "Pakistan_AI_Policy_Monitor_Evidence_Register_and_Coverage_Log_"
    "2026-08-18_to_2026-09-03.xlsx"
)

EXPECTED_SHEETS = [
    "Read Me",
    "Evidence Register",
    "Coverage & Gaps",
    "Exclusions",
    "Briefing Units",
    "Sources",
    "Watchlist",
    "File Inventory",
    "Methodology",
]

FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def read_csv(name: str) -> list[dict[str, str]]:
    """Read one audited CSV input."""
    with (INTERMEDIATE / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def check(condition: bool, message: str) -> None:
    """Raise a verification error when a requirement is not satisfied."""
    if not condition:
        raise AssertionError(message)


def sheet_records(worksheet: Any, columns: int) -> list[list[Any]]:
    """Read populated table rows beginning at workbook row five."""
    return [
        [worksheet.cell(row=row, column=column).value for column in range(1, columns + 1)]
        for row in range(5, worksheet.max_row + 1)
        if worksheet.cell(row=row, column=1).value not in (None, "")
    ]


def normalized_value(value: Any, field: str) -> str:
    """Normalize a typed cell to the source CSV's lexical form."""
    if value is None:
        return ""
    if isinstance(value, bool):
        if field in {"readable", "material_gap"}:
            return "yes" if value else "no"
        return "true" if value else "false"
    if isinstance(value, datetime):
        if field in {
            "event_date",
            "publication_date",
            "first_detected",
            "last_updated",
            "date_or_trigger",
            "nominal_date",
            "access_date_utc",
        } and value.time() == datetime.min.time():
            return value.date().isoformat()
        base = value.strftime("%Y-%m-%dT%H:%M:%S")
        if field.endswith("_pkt"):
            return f"{base}+05:00"
        return f"{base}Z"
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def compare_sheet_to_csv(
    worksheet: Any,
    csv_name: str,
    *,
    extra_fields: dict[str, Callable[[dict[str, str]], str]] | None = None,
    start_row: int = 5,
) -> None:
    """Compare every populated cell with its audited CSV source."""
    source = read_csv(csv_name)
    fields = list(source[0]) if source else []
    extra_fields = extra_fields or {}
    fields.extend(extra_fields)
    actual = [
        [
            worksheet.cell(row=row, column=column).value
            for column in range(1, len(fields) + 1)
        ]
        for row in range(start_row, start_row + len(source))
    ]
    check(len(actual) == len(source), f"{worksheet.title}: source-row count mismatch")
    for row_index, (source_row, workbook_row) in enumerate(
        zip(source, actual), start=start_row
    ):
        for column_index, field in enumerate(fields):
            expected = (
                extra_fields[field](source_row)
                if field in extra_fields
                else source_row.get(field, "").strip()
            )
            observed = normalized_value(workbook_row[column_index], field)
            if field in {
                "gap_minutes",
                "gap_hours",
            }:
                check(
                    math.isclose(float(observed), float(expected), rel_tol=0.0, abs_tol=1e-9),
                    f"{worksheet.title}!{row_index},{column_index + 1}: numeric mismatch",
                )
            elif field in {
                "window_start_utc",
                "window_end_utc",
                "window_start_pkt",
                "window_end_pkt",
                "filesystem_birthtime",
                "filesystem_mtime",
                "internal_created",
                "internal_modified",
            } and expected:
                check(
                    observed[:19] == expected[:19],
                    f"{worksheet.title}!{row_index},{column_index + 1}: datetime mismatch",
                )
            else:
                check(
                    observed == expected,
                    (
                        f"{worksheet.title}!{row_index},{column_index + 1}: "
                        f"expected {expected!r}, observed {observed!r}"
                    ),
                )


def all_formulas(workbook: Any) -> list[tuple[str, str, str]]:
    """Collect every formula coordinate and expression."""
    formulas: list[tuple[str, str, str]] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((worksheet.title, cell.coordinate, cell.value))
    return formulas


def canonical_formula(formula: str) -> str:
    """Normalize Excel's optional quotes around simple worksheet names."""
    return re.sub(r"'([A-Za-z0-9_]+)'!", r"\1!", formula)


def verify_workbook(require_cached_formulas: bool = False) -> list[str]:
    """Run structural, data, formula, source, and presentation checks."""
    results: list[str] = []
    check(WORKBOOK_PATH.is_file(), f"Workbook not found: {WORKBOOK_PATH}")
    check(WORKBOOK_PATH.stat().st_size > 50_000, "Workbook is unexpectedly small")
    xlsx_files = list(OUTPUT_DIR.glob("*.xlsx"))
    check(xlsx_files == [WORKBOOK_PATH], "Output directory must contain exactly one XLSX")
    check(not WORKBOOK_PATH.with_suffix(".building").exists(), "Temporary workbook remains")
    results.append("single final XLSX exists and no build temporary remains")

    with zipfile.ZipFile(WORKBOOK_PATH) as archive:
        check(archive.testzip() is None, "XLSX ZIP CRC failure")
        names = set(archive.namelist())
        package_parts = {name: archive.read(name) for name in names}
        check("xl/workbook.xml" in names, "Missing workbook XML")
        check("docProps/core.xml" in names, "Missing core properties")
        check(not any(name.startswith("xl/externalLinks/") for name in names), "External links present")
        check(not any("vbaProject" in name for name in names), "Unexpected macros present")
        check(
            not any(b"/Users/" in payload or b":\\Users\\" in payload for payload in package_parts.values()),
            "Absolute local filesystem path is embedded in the workbook package",
        )
        check(
            b"x15ac:absPath" not in package_parts["xl/workbook.xml"],
            "Native Excel absolute-path metadata is present",
        )
        worksheet_xml = [
            package_parts[name]
            for name in names
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        ]
        table_xml = [
            package_parts[name]
            for name in names
            if name.startswith("xl/tables/table") and name.endswith(".xml")
        ]
        check(
            all(b"<autoFilter" not in xml for xml in worksheet_xml),
            "Worksheet-level AutoFilter duplicates a table filter",
        )
        check(len(table_xml) == 7, "Expected seven table definition parts")
        check(
            all(b"<autoFilter" in xml for xml in table_xml),
            "A table definition is missing its AutoFilter",
        )
    results.append(
        "OOXML package is intact, path-clean, macro-free, externally unlinked, and uses table-only filters"
    )

    workbook = load_workbook(WORKBOOK_PATH, data_only=False)
    check(workbook.sheetnames == EXPECTED_SHEETS, "Worksheet order/name mismatch")
    check(all(ws.sheet_state == "visible" for ws in workbook.worksheets), "Hidden sheet found")
    check(workbook.index(workbook.active) == 0, "Read Me is not the active sheet")
    check(workbook.properties.creator == "Muhammad Umar Zafar", "Creator metadata mismatch")
    check(
        workbook.calculation.calcMode in (None, "auto"),
        "Workbook calculation mode is explicitly non-automatic",
    )
    check(workbook.calculation.fullCalcOnLoad is True, "Full calculation on load is not enabled")
    check(workbook.calculation.forceFullCalc is True, "Forced full calculation is not enabled")
    results.append(
        "nine visible sheets, metadata, active sheet, and calculation settings pass"
    )

    expected_shapes = {
        "Read Me": (45, 9),
        "Evidence Register": (56, 26),
        "Coverage & Gaps": (41, 8),
        "Exclusions": (17, 5),
        "Briefing Units": (18, 11),
        "Sources": (115, 10),
        "Watchlist": (18, 6),
        "File Inventory": (17, 17),
        "Methodology": (48, 8),
    }
    for sheet_name, (rows, columns) in expected_shapes.items():
        worksheet = workbook[sheet_name]
        check(
            (worksheet.max_row, worksheet.max_column) == (rows, columns),
            f"{sheet_name}: expected {rows}x{columns}, got {worksheet.max_row}x{worksheet.max_column}",
        )
        check(worksheet["A1"].value not in (None, ""), f"{sheet_name}: missing title")
        check(worksheet.sheet_view.showGridLines is False, f"{sheet_name}: gridlines visible")
        check(worksheet.freeze_panes is not None, f"{sheet_name}: freeze pane missing")
        check(worksheet.print_area, f"{sheet_name}: print area missing")
    results.append("all worksheet shapes, titles, print areas, gridline settings, and freeze panes pass")

    table_expectations = {
        "Evidence Register": ("tblEvidenceRegister", "A4:Z56"),
        "Coverage & Gaps": ("tblCoverageGaps", "A28:H41"),
        "Exclusions": ("tblExclusions", "A4:E17"),
        "Briefing Units": ("tblBriefingUnits", "A4:K18"),
        "Sources": ("tblSources", "A4:J115"),
        "Watchlist": ("tblWatchlist", "A4:F18"),
        "File Inventory": ("tblFileInventory", "A4:Q17"),
    }
    for sheet_name, (table_name, ref) in table_expectations.items():
        tables = workbook[sheet_name].tables
        check(list(tables) == [table_name], f"{sheet_name}: table name mismatch")
        check(tables[table_name].ref == ref, f"{sheet_name}: table range mismatch")
    results.append("all seven data registers have the expected filterable Excel tables")

    inventory_rows = read_csv("dataset_inventory.csv")
    inventory_hashes = {row["source_file"]: row["sha256"] for row in inventory_rows}
    compare_sheet_to_csv(workbook["Evidence Register"], "evidence_register.csv")
    compare_sheet_to_csv(workbook["Exclusions"], "exclusion_log.csv")
    compare_sheet_to_csv(workbook["Sources"], "source_register.csv")
    compare_sheet_to_csv(workbook["Watchlist"], "watchlist.csv")
    compare_sheet_to_csv(workbook["File Inventory"], "dataset_inventory.csv")
    compare_sheet_to_csv(
        workbook["Coverage & Gaps"], "coverage_gaps.csv", start_row=29
    )
    compare_sheet_to_csv(
        workbook["Briefing Units"],
        "briefing_units.csv",
        extra_fields={"source_file_sha256": lambda row: inventory_hashes[row["source_file"]]},
    )
    results.append("every data-register cell reconciles to its audited CSV input")

    evidence_ws = workbook["Evidence Register"]
    evidence_rows = sheet_records(evidence_ws, 26)
    evidence_ids = [row[0] for row in evidence_rows]
    lifecycles = [row[25] for row in evidence_rows]
    confidences = [row[21] for row in evidence_rows]
    check(len(evidence_ids) == len(set(evidence_ids)) == 52, "Evidence IDs are not unique")
    check(sum(str(value).startswith("EVT-") for value in evidence_ids) == 36, "Event count mismatch")
    check(sum(str(value).startswith("BASE-") for value in evidence_ids) == 16, "Baseline count mismatch")
    check(Counter(confidences) == Counter({"High": 32, "Medium": 16, "Low": 4}), "Confidence totals mismatch")
    check(
        Counter(lifecycles)
        == Counter(
            {
                "IN_SPAN_DEVELOPMENT": 35,
                "BASELINE_CARRY_IN": 13,
                "CONTRADICTED_OR_UNRESOLVED": 4,
            }
        ),
        "Lifecycle totals mismatch",
    )
    disputed_baselines = sum(
        str(event_id).startswith("BASE-") and lifecycle == "CONTRADICTED_OR_UNRESOLVED"
        for event_id, lifecycle in zip(evidence_ids, lifecycles)
    )
    check(disputed_baselines == 3, "Disputed baseline count mismatch")
    results.append("event/baseline, confidence, and lifecycle controls independently reconcile")

    read_me = workbook["Read Me"]
    expected_live_formulas = {
        13: "=COUNTA('Evidence Register'!$A$5:$A$56)",
        14: '=COUNTIF(\'Evidence Register\'!$A$5:$A$56,"EVT-*")',
        15: '=COUNTIF(\'Evidence Register\'!$A$5:$A$56,"BASE-*")',
        16: '=COUNTIFS(\'Evidence Register\'!$A$5:$A$56,"BASE-*",\'Evidence Register\'!$Z$5:$Z$56,"BASELINE_CARRY_IN")',
        17: '=COUNTIFS(\'Evidence Register\'!$A$5:$A$56,"BASE-*",\'Evidence Register\'!$Z$5:$Z$56,"CONTRADICTED_OR_UNRESOLVED")',
        18: "=COUNTA('Sources'!$A$5:$A$115)",
        19: "=COUNTA('Briefing Units'!$A$5:$A$18)",
        20: "=COUNTA('Exclusions'!$A$5:$A$17)",
        21: "=COUNTA('Watchlist'!$A$5:$A$18)",
        22: '=COUNTIF(\'Evidence Register\'!$V$5:$V$56,"High")',
        23: '=COUNTIF(\'Evidence Register\'!$V$5:$V$56,"Medium")',
        24: '=COUNTIF(\'Evidence Register\'!$V$5:$V$56,"Low")',
    }
    expected_values = [52, 36, 16, 13, 3, 111, 14, 13, 14, 32, 16, 4]
    for row_number, expected in expected_live_formulas.items():
        check(
            canonical_formula(read_me.cell(row=row_number, column=2).value)
            == canonical_formula(expected),
            f"Read Me!B{row_number} formula mismatch",
        )
        check(read_me.cell(row=row_number, column=2).comment is not None, f"Read Me!B{row_number} formula comment missing")
        check(read_me.cell(row=row_number, column=5).value == f'=IF(B{row_number}=D{row_number},"PASS","CHECK")', f"Read Me!E{row_number} reconciliation formula mismatch")
    check([read_me.cell(row=row, column=4).value for row in range(13, 25)] == expected_values, "Expected control totals mismatch")
    results.append("Read Me formulas, expected totals, reconciliation formulas, and comments pass")

    coverage = workbook["Coverage & Gaps"]
    expected_coverage = {
        "B6": 1440,
        "B7": "=COUNTA('Briefing Units'!$A$5:$A$18)",
        "B8": "=B6-B7",
        "B9": "=IF(B6=0,0,B7/B6)",
        "B10": 22017,
        "B11": 215.82,
        "B12": "=IF(B10=0,0,B11/B10)",
        "B13": 1468,
        "B14": "=B13-B7",
        "B15": "=IF(B13=0,0,B7/B13)",
    }
    for coordinate, expected in expected_coverage.items():
        observed = coverage[coordinate].value
        formulas_match = (
            isinstance(observed, str)
            and isinstance(expected, str)
            and canonical_formula(observed) == canonical_formula(expected)
        )
        check(observed == expected or formulas_match, f"Coverage & Gaps!{coordinate} mismatch")
        check(coverage[coordinate].comment is not None, f"Coverage & Gaps!{coordinate} comment missing")
    gap_rows = [
        [coverage.cell(row=row, column=column).value for column in range(1, 9)]
        for row in range(29, 42)
    ]
    check(len(gap_rows) == 13, "Gap-row count mismatch")
    check(math.isclose(sum(row[4] for row in gap_rows), 21801.18, abs_tol=1e-7), "Gap-minute total mismatch")
    check(sum(row[6] for row in gap_rows) == 1454, "Gap missing-run total mismatch")
    check(all(row[7] is True for row in gap_rows), "Non-material gap found")
    check(math.isclose(14 / 1440, 0.009722222222222222, abs_tol=1e-15), "Nominal share arithmetic mismatch")
    check(math.isclose(215.82 / 22017, 0.009802425398555662, abs_tol=1e-15), "Strict-window share arithmetic mismatch")
    results.append("coverage inputs, formulas, gap totals, and divide-by-zero guards pass")

    source_ws = workbook["Sources"]
    source_ids = [row[0] for row in sheet_records(source_ws, 10)]
    check(len(source_ids) == len(set(source_ids)) == 111, "Source IDs are not unique")
    for worksheet, url_columns in (
        (evidence_ws, (12, 14)),
        (source_ws, (4,)),
        (workbook["Watchlist"], (6,)),
    ):
        for row in range(5, worksheet.max_row + 1):
            for column in url_columns:
                cell = worksheet.cell(row=row, column=column)
                if cell.value in (None, ""):
                    continue
                urls = re.findall(r"https?://[^\s|]+", str(cell.value))
                check(urls, f"{worksheet.title}!{cell.coordinate}: no visible URL")
                check(all(url.startswith(("http://", "https://")) for url in urls), f"{worksheet.title}!{cell.coordinate}: malformed URL")
                if len(re.findall(r"https?://", str(cell.value))) == 1 and re.fullmatch(r"https?://\S+", str(cell.value)):
                    check(cell.hyperlink is not None, f"{worksheet.title}!{cell.coordinate}: single URL is not clickable")
                else:
                    check(cell.hyperlink is None, f"{worksheet.title}!{cell.coordinate}: multi-URL cell has an invalid combined hyperlink")
    for row in range(5, evidence_ws.max_row + 1):
        check(
            evidence_ws.cell(row=row, column=12).value
            or evidence_ws.cell(row=row, column=14).value,
            f"Evidence Register row {row}: no exposed source URL",
        )
    results.append("visible source URLs, hyperlink behavior, and evidence-row URL coverage pass")

    def find_row(worksheet: Any, identifier: str) -> int:
        """Find an identifier in column A."""
        for row in range(5, worksheet.max_row + 1):
            if worksheet.cell(row=row, column=1).value == identifier:
                return row
        raise AssertionError(f"{worksheet.title}: missing {identifier}")

    who_row = find_row(evidence_ws, "EVT-PROC-20260828-WHO-BALOCHISTAN-GIS")
    who_pub = str(evidence_ws.cell(row=who_row, column=7).value)
    check("24" in who_pub and "28" in who_pub, "WHO evidence date conflict not preserved")
    check(evidence_ws.cell(row=who_row, column=7).comment is not None, "WHO evidence conflict comment missing")
    s047_row = find_row(source_ws, "S047")
    s047_pub = str(source_ws.cell(row=s047_row, column=6).value)
    s047_limits = str(source_ws.cell(row=s047_row, column=10).value)
    check("24" in s047_pub and "28" in s047_pub, "S047 publication conflict not preserved")
    check("24" in s047_limits and "28" in s047_limits, "S047 limitation conflict not preserved")
    telecom_row = find_row(evidence_ws, "EVT-TEL-20260828-AMENDMENT-WITHDRAWN")
    telecom_pub = evidence_ws.cell(row=telecom_row, column=7).value
    telecom_date = telecom_pub.date() if isinstance(telecom_pub, datetime) else telecom_pub
    check(
        isinstance(telecom_date, date) and telecom_date.isoformat() == "2026-08-28",
        "Telecom publication date mismatch",
    )
    results.append("WHO mutable-date conflict and telecom publication-date spot checks pass")

    briefing_ws = workbook["Briefing Units"]
    briefing_rows = sheet_records(briefing_ws, 11)
    check(len(briefing_rows) == 14, "Briefing-unit count mismatch")
    check(all(isinstance(row[4], datetime) and isinstance(row[6], datetime) for row in briefing_rows), "Briefing datetimes are not typed")
    check(all(re.fullmatch(r"[0-9a-f]{64}", str(row[10])) for row in briefing_rows), "Briefing hash mismatch")
    check(len(set(row[10] for row in briefing_rows[:3])) == 1, "B001–B003 should share the source-file hash")
    inventory_ws = workbook["File Inventory"]
    inventory_records = sheet_records(inventory_ws, 17)
    check(all(isinstance(row[3], int) for row in inventory_records), "Inventory bytes are not typed integers")
    check(all(isinstance(row[13], int) for row in inventory_records), "Inventory briefing-unit counts are not typed integers")
    check(all(isinstance(row[7], bool) for row in inventory_records), "Inventory readability is not typed Boolean")
    results.append("briefing UTC/PKT values, shared hashes, and inventory numeric types pass")

    formulas = all_formulas(workbook)
    check(formulas, "Workbook contains no formulas")
    for sheet_name, coordinate, formula in formulas:
        check(not any(error in formula for error in FORMULA_ERRORS), f"{sheet_name}!{coordinate}: formula error token")
        check("Evidence Register!" not in formula, f"{sheet_name}!{coordinate}: unquoted cross-sheet reference")
        check("Coverage & Gaps!" not in formula, f"{sheet_name}!{coordinate}: unquoted cross-sheet reference")
        check("Briefing Units!" not in formula, f"{sheet_name}!{coordinate}: unquoted cross-sheet reference")
    results.append(f"{len(formulas)} formulas are present with no error tokens or unquoted spaced-sheet references")

    check(len(evidence_ws.data_validations.dataValidation) == 2, "Evidence data validations missing")
    check(len(evidence_ws.conditional_formatting) >= 2, "Evidence conditional formatting missing")
    check(len(coverage.conditional_formatting) >= 2, "Coverage conditional formatting missing")
    check(len(inventory_ws.conditional_formatting) >= 2, "Inventory conditional formatting missing")
    check(evidence_ws.column_dimensions["P"].width >= 50, "Evidence facts column too narrow")
    check(source_ws.column_dimensions["D"].width >= 55, "Source URL column too narrow")
    check(max(evidence_ws.row_dimensions[row].height or 0 for row in range(5, 57)) >= 100, "Evidence long-text rows were not expanded")
    results.append("validation rules, conditional formats, widths, and long-text row heights pass")

    cached = load_workbook(WORKBOOK_PATH, data_only=True)
    cached_coordinates = {
        "Read Me": {
            "B13": 52,
            "B14": 36,
            "B15": 16,
            "B16": 13,
            "B17": 3,
            "B18": 111,
            "B19": 14,
            "B20": 13,
            "B21": 14,
            "B22": 32,
            "B23": 16,
            "B24": 4,
        },
        "Coverage & Gaps": {
            "B7": 14,
            "B8": 1426,
            "B9": 14 / 1440,
            "B12": 215.82 / 22017,
            "B14": 1454,
            "B15": 14 / 1468,
        },
    }
    cached_values = [
        cached[sheet][coordinate].value
        for sheet, coordinates in cached_coordinates.items()
        for coordinate in coordinates
    ]
    if require_cached_formulas:
        for sheet, coordinates in cached_coordinates.items():
            for coordinate, expected in coordinates.items():
                observed = cached[sheet][coordinate].value
                check(observed is not None, f"{sheet}!{coordinate}: cached formula value missing")
                if isinstance(expected, float):
                    check(
                        math.isclose(float(observed), expected, rel_tol=1e-12, abs_tol=1e-12),
                        f"{sheet}!{coordinate}: cached value mismatch",
                    )
                else:
                    check(observed == expected, f"{sheet}!{coordinate}: cached value mismatch")
        for row in range(13, 25):
            check(cached["Read Me"].cell(row=row, column=5).value == "PASS", f"Read Me!E{row}: cached reconciliation is not PASS")
        results.append("Excel-cached formula values and all twelve PASS reconciliations match")
    elif all(value is None for value in cached_values):
        results.append("formula structures pass; cached values await native Excel recalculation")
    else:
        results.append("formula structures pass and cached values are present")

    for worksheet in cached.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    check(
                        cell.value not in FORMULA_ERRORS,
                        f"{worksheet.title}!{cell.coordinate}: cached formula error {cell.value}",
                    )
    results.append("no cached Excel error values found")
    return results


def main() -> None:
    """Run verification and print a compact gate report."""
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--require-cached-formulas",
        action="store_true",
        help="Require formula results written by native Excel after recalculation/save.",
    )
    args = parser.parse_args()
    results = verify_workbook(require_cached_formulas=args.require_cached_formulas)
    print("VERIFY WORKBOOK")
    print("===============")
    for index, result in enumerate(results, start=1):
        print(f"PASS {index:02d}: {result}")
    print(f"PASS: {len(results)}/{len(results)} gates")


if __name__ == "__main__":
    main()
