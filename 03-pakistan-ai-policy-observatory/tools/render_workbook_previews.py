#!/usr/bin/env python3
"""Render compact visual-QA previews of every workbook sheet with OpenPyXL."""

from __future__ import annotations

import math
import textwrap
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont


__all__ = ["render_all"]

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "outputs" / "01a0653c-96d7-7c92-918b-3975d6716164" / (
    "Pakistan_AI_Policy_Monitor_Evidence_Register_and_Coverage_Log_"
    "2026-08-18_to_2026-09-03.xlsx"
)
OUTPUT_DIR = ROOT / "intermediate" / "workbook_visual_qa"

RANGES: dict[str, list[tuple[int, int, int, int]]] = {
    "Read Me": [(1, 45, 1, 9)],
    "Evidence Register": [
        (1, 16, 1, 14),
        (1, 16, 15, 26),
        (20, 26, 1, 14),
        (20, 26, 15, 26),
        (50, 56, 1, 14),
        (50, 56, 15, 26),
    ],
    "Coverage & Gaps": [(1, 41, 1, 8)],
    "Exclusions": [(1, 17, 1, 5)],
    "Briefing Units": [(1, 18, 1, 11)],
    "Sources": [(1, 14, 1, 10), (46, 56, 1, 10), (106, 115, 1, 10)],
    "Watchlist": [(1, 18, 1, 6)],
    "File Inventory": [(1, 17, 1, 9), (1, 17, 10, 17)],
    "Methodology": [(1, 25, 1, 8), (26, 48, 1, 8)],
}

FORMULA_DISPLAYS: dict[tuple[str, str], Any] = {
    **{
        ("Read Me", f"B{row}"): value
        for row, value in zip(
            range(13, 25), [52, 36, 16, 13, 3, 111, 14, 13, 14, 32, 16, 4]
        )
    },
    **{("Read Me", f"E{row}"): "PASS" for row in range(13, 25)},
    ("Read Me", "H13"): 1440,
    ("Read Me", "H14"): 14,
    ("Read Me", "H15"): 1426,
    ("Read Me", "H16"): 14 / 1440,
    ("Read Me", "H17"): 215.82 / 22017,
    ("Coverage & Gaps", "B7"): 14,
    ("Coverage & Gaps", "B8"): 1426,
    ("Coverage & Gaps", "B9"): 14 / 1440,
    ("Coverage & Gaps", "B12"): 215.82 / 22017,
    ("Coverage & Gaps", "B14"): 1454,
    ("Coverage & Gaps", "B15"): 14 / 1468,
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    """Load a portable system font for preview rendering."""
    candidates = (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    )
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def rgb(color: Any, default: str = "FFFFFF") -> tuple[int, int, int]:
    """Convert an OpenPyXL RGB color to a Pillow color."""
    value = getattr(color, "rgb", None)
    if not isinstance(value, str) or len(value) not in {6, 8}:
        value = default
    value = value[-6:]
    try:
        return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        return tuple(int(default[index : index + 2], 16) for index in (0, 2, 4))


def column_pixels(worksheet: Any, column: int) -> int:
    """Approximate Excel column width in pixels."""
    width = worksheet.column_dimensions[get_column_letter(column)].width or 8.43
    return max(28, int(width * 7.0 + 5))


def row_pixels(worksheet: Any, row: int) -> int:
    """Approximate Excel row height in pixels."""
    height = worksheet.row_dimensions[row].height or 15.0
    return max(20, int(height * 96 / 72))


def display_value(sheet_name: str, cell: Any) -> str:
    """Return the workbook value as it should appear in a QA preview."""
    value = FORMULA_DISPLAYS.get((sheet_name, cell.coordinate), cell.value)
    if value is None:
        return ""
    if isinstance(value, datetime):
        suffix = " PKT" if "PKT" in (cell.number_format or "") else " UTC"
        return value.strftime("%Y-%m-%d %H:%M") + suffix
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and "%" in (cell.number_format or ""):
        return f"{value:.2%}"
    if isinstance(value, (int, float)):
        if cell.number_format == "#,##0.00":
            return f"{value:,.2f}"
        if cell.number_format == "#,##0.0":
            return f"{value:,.1f}"
        if "#,##0" in (cell.number_format or ""):
            return f"{value:,.0f}"
    if isinstance(value, str) and value.startswith("="):
        return "[formula]"
    return str(value)


def table_fill(worksheet: Any, row: int, column: int) -> tuple[int, int, int] | None:
    """Approximate table striping that is applied by Excel at render time."""
    for table in worksheet.tables.values():
        min_column, min_row, max_column, max_row = range_boundaries(table.ref)
        if min_column <= column <= max_column and min_row < row <= max_row:
            return (233, 239, 245) if (row - min_row) % 2 == 1 else (255, 255, 255)
    return None


def merged_ranges(worksheet: Any, bounds: tuple[int, int, int, int]) -> list[Any]:
    """Return merged ranges that intersect a preview tile."""
    row_start, row_end, column_start, column_end = bounds
    result = []
    for merged in worksheet.merged_cells.ranges:
        if not (
            merged.max_row < row_start
            or merged.min_row > row_end
            or merged.max_col < column_start
            or merged.min_col > column_end
        ):
            result.append(merged)
    return result


def draw_wrapped_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    box: tuple[int, int, int, int],
    *,
    text_color: tuple[int, int, int],
    text_font: ImageFont.FreeTypeFont | ImageFont.ImageFont,
    align: str,
) -> None:
    """Draw clipped, wrapped preview text inside a cell rectangle."""
    left, top, right, bottom = box
    width = max(8, right - left - 8)
    average_character = max(4, int(getattr(text_font, "size", 9) * 0.56))
    characters = max(1, width // average_character)
    paragraphs = text.splitlines() or [""]
    lines: list[str] = []
    for paragraph in paragraphs:
        lines.extend(textwrap.wrap(paragraph, width=characters, break_long_words=True) or [""])
    line_height = max(10, int(getattr(text_font, "size", 9) * 1.25))
    max_lines = max(1, (bottom - top - 6) // line_height)
    lines = lines[:max_lines]
    y = top + 3
    for line in lines:
        line_width = draw.textlength(line, font=text_font)
        if align == "right":
            x = max(left + 4, right - 4 - line_width)
        elif align == "center":
            x = max(left + 4, left + (right - left - line_width) / 2)
        else:
            x = left + 4
        draw.text((x, y), line, fill=text_color, font=text_font)
        y += line_height


def render_tile(
    worksheet: Any,
    bounds: tuple[int, int, int, int],
    tile_number: int,
) -> Image.Image:
    """Render one sheet-range tile with widths, heights, fills, and text."""
    row_start, row_end, column_start, column_end = bounds
    widths = [column_pixels(worksheet, column) for column in range(column_start, column_end + 1)]
    heights = [row_pixels(worksheet, row) for row in range(row_start, row_end + 1)]
    label_height = 34
    image = Image.new(
        "RGB", (sum(widths) + 2, sum(heights) + label_height + 2), (255, 255, 255)
    )
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, image.width, label_height), fill=(16, 35, 52))
    label = (
        f"{worksheet.title} — tile {tile_number}: "
        f"{get_column_letter(column_start)}{row_start}:"
        f"{get_column_letter(column_end)}{row_end}"
    )
    draw.text((10, 8), label, fill=(255, 255, 255), font=font(13, bold=True))

    x_positions = [1]
    for width in widths:
        x_positions.append(x_positions[-1] + width)
    y_positions = [label_height + 1]
    for height in heights:
        y_positions.append(y_positions[-1] + height)

    merges = merged_ranges(worksheet, bounds)
    merged_cells = {
        (row, column)
        for merged in merges
        for row in range(max(row_start, merged.min_row), min(row_end, merged.max_row) + 1)
        for column in range(max(column_start, merged.min_col), min(column_end, merged.max_col) + 1)
        if (row, column) != (merged.min_row, merged.min_col)
    }

    for row in range(row_start, row_end + 1):
        for column in range(column_start, column_end + 1):
            if (row, column) in merged_cells:
                continue
            cell = worksheet.cell(row=row, column=column)
            x1 = x_positions[column - column_start]
            x2 = x_positions[column - column_start + 1]
            y1 = y_positions[row - row_start]
            y2 = y_positions[row - row_start + 1]

            merged = next(
                (
                    item
                    for item in merges
                    if item.min_row == row and item.min_col == column
                ),
                None,
            )
            if merged is not None:
                clipped_max_column = min(column_end, merged.max_col)
                clipped_max_row = min(row_end, merged.max_row)
                x2 = x_positions[clipped_max_column - column_start + 1]
                y2 = y_positions[clipped_max_row - row_start + 1]

            fill_color = rgb(cell.fill.fgColor)
            if cell.fill.fill_type is None:
                fill_color = table_fill(worksheet, row, column) or (255, 255, 255)
            draw.rectangle((x1, y1, x2, y2), fill=fill_color, outline=(205, 214, 219), width=1)

            value = display_value(worksheet.title, cell)
            if not value or isinstance(cell, MergedCell):
                continue
            size = 13 if row == 1 else 8 if row == 2 else 9
            is_bold = bool(cell.font.bold) or row in {1, 4, 5, 12, 18, 27, 28}
            text_font = font(size, bold=is_bold)
            text_color = rgb(cell.font.color, default="22313B")
            horizontal = cell.alignment.horizontal or "left"
            draw_wrapped_text(
                draw,
                value,
                (x1, y1, x2, y2),
                text_color=text_color,
                text_font=text_font,
                align=horizontal,
            )

    if image.width > 2600:
        ratio = 2600 / image.width
        image = image.resize(
            (2600, max(1, int(image.height * ratio))),
            resample=Image.Resampling.LANCZOS,
        )
    return image


def render_sheet(worksheet: Any, ranges: list[tuple[int, int, int, int]]) -> Path:
    """Stack all QA tiles for one worksheet into a single PNG."""
    tiles = [render_tile(worksheet, bounds, index) for index, bounds in enumerate(ranges, start=1)]
    width = max(tile.width for tile in tiles)
    spacing = 24
    height = sum(tile.height for tile in tiles) + spacing * (len(tiles) - 1)
    canvas = Image.new("RGB", (width, height), (235, 239, 241))
    y = 0
    for tile in tiles:
        canvas.paste(tile, (0, y))
        y += tile.height + spacing
    output = OUTPUT_DIR / f"{worksheet.title.lower().replace(' ', '_').replace('&', 'and')}.png"
    canvas.save(output, format="PNG", optimize=True)
    return output


def render_all() -> list[Path]:
    """Render a QA preview for every sheet in the final workbook."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(WORKBOOK_PATH, data_only=False)
    outputs = [render_sheet(workbook[name], RANGES[name]) for name in workbook.sheetnames]
    return outputs


def main() -> None:
    """Render and list all sheet-preview paths."""
    for output in render_all():
        print(output)


if __name__ == "__main__":
    main()
